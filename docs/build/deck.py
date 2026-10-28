"""Une description de page, deux rendus : PPTX (texte éditable) et PNG (pour le PDF, avec contrôle des débordements)."""
from dataclasses import dataclass, field

from PIL import Image as PImage, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

W_IN, H_IN = 13.333, 7.5
PX = 144  # px par pouce pour le rendu PNG (1920 × 1080)
FONT = "Arial"
F_REG, F_BOLD, F_IT = "C:/Windows/Fonts/arial.ttf", "C:/Windows/Fonts/arialbd.ttf", "C:/Windows/Fonts/ariali.ttf"
BLUE = (31, 78, 160)
BORDEAUX = (128, 20, 40)
DARK = (30, 30, 30)
GREY = (110, 110, 110)
LIGHT = (240, 243, 248)


@dataclass
class Run:
    text: str
    bold: bool = False
    italic: bool = False
    color: tuple = DARK


@dataclass
class Para:
    runs: list
    size: float = 14
    bullet: bool = False
    space_after: float = 6


@dataclass
class Text:
    x: float
    y: float
    w: float
    h: float
    paras: list
    fill: tuple = None
    line: tuple = None
    anchor_middle: bool = False


@dataclass
class Img:
    x: float
    y: float
    w: float
    h: float
    path: str


@dataclass
class Rect:
    x: float
    y: float
    w: float
    h: float
    fill: tuple = None
    line: tuple = None


@dataclass
class Table:
    x: float
    y: float
    widths: list
    rows: list  # liste de listes de str ; la première est l'en-tête
    size: float = 12
    row_h: float = 0.36


@dataclass
class Slide:
    title: str = ""
    items: list = field(default_factory=list)


def P(text, size=14, bold=False, color=DARK, bullet=False, italic=False, after=6):
    return Para([Run(text, bold, italic, color)], size, bullet, after)


def RP(label, text, size=14, after=6):
    """Rubrique en gras suivie de son texte."""
    return Para([Run(label + " ", True), Run(text)], size, False, after)


# ---------------------------------------------------------------- PPTX
def _rgb(c):
    return RGBColor(*c)


def _jpg(path):
    """Copie JPEG d'une image, pour garder le PPTX léger (import dans Google Slides)."""
    from pathlib import Path
    src = Path(path)
    dst = Path(__file__).parent / "jpg" / (src.stem + ".jpg")
    if not dst.exists() or dst.stat().st_mtime < src.stat().st_mtime:
        dst.parent.mkdir(exist_ok=True)
        im = PImage.open(src).convert("RGB")
        if im.width > 2000 and not src.stem.startswith("sitac"):
            im = im.resize((2000, round(im.height * 2000 / im.width)), PImage.LANCZOS)
        im.save(dst, quality=82, optimize=True)
    return str(dst)


def to_pptx(slides, path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W_IN), Inches(H_IN)
    blank = prs.slide_layouts[6]
    for s in slides:
        sl = prs.slides.add_slide(blank)
        if s.title:
            _text(sl, Text(0.5, 0.3, 12.3, 0.7, [Para([Run(s.title, True)], 26, space_after=0)]))
        for it in s.items:
            if isinstance(it, Text):
                _text(sl, it)
            elif isinstance(it, Img):
                sl.shapes.add_picture(_jpg(it.path), Inches(it.x), Inches(it.y), Inches(it.w), Inches(it.h))
            elif isinstance(it, Rect):
                sh = sl.shapes.add_shape(1, Inches(it.x), Inches(it.y), Inches(it.w), Inches(it.h))
                if it.fill:
                    sh.fill.solid()
                    sh.fill.fore_color.rgb = _rgb(it.fill)
                else:
                    sh.fill.background()
                if it.line:
                    sh.line.color.rgb = _rgb(it.line)
                else:
                    sh.line.fill.background()
                sh.shadow.inherit = False
            elif isinstance(it, Table):
                _table(sl, it)
    prs.save(path)


def _text(sl, t):
    tb = sl.shapes.add_textbox(Inches(t.x), Inches(t.y), Inches(t.w), Inches(t.h))
    if t.fill:
        tb.fill.solid()
        tb.fill.fore_color.rgb = _rgb(t.fill)
    if t.line:
        tb.line.color.rgb = _rgb(t.line)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.08)
    tf.margin_top = tf.margin_bottom = Inches(0.05)
    if t.anchor_middle:
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, p in enumerate(t.paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = PP_ALIGN.LEFT
        para.space_after = Pt(p.space_after)
        runs = ([Run("• ")] + p.runs) if p.bullet else p.runs
        for r in runs:
            rr = para.add_run()
            rr.text = r.text
            rr.font.name = FONT
            rr.font.size = Pt(p.size)
            rr.font.bold = r.bold
            rr.font.italic = r.italic
            rr.font.color.rgb = _rgb(r.color)


def _table(sl, t):
    rows, cols = len(t.rows), len(t.widths)
    shape = sl.shapes.add_table(rows, cols, Inches(t.x), Inches(t.y), Inches(sum(t.widths)), Inches(t.row_h * rows))
    tbl = shape.table
    for j, w in enumerate(t.widths):
        tbl.columns[j].width = Inches(w)
    for i, row in enumerate(t.rows):
        tbl.rows[i].height = Inches(t.row_h)
        for j, val in enumerate(row):
            cell = tbl.cell(i, j)
            cell.margin_left = cell.margin_right = Inches(0.06)
            cell.margin_top = cell.margin_bottom = Inches(0.02)
            cell.fill.solid()
            cell.fill.fore_color.rgb = _rgb(BLUE if i == 0 else ((255, 255, 255) if i % 2 else LIGHT))
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT
            r = p.add_run()
            r.text = val
            r.font.name = FONT
            r.font.size = Pt(t.size)
            r.font.bold = i == 0
            r.font.color.rgb = _rgb((255, 255, 255) if i == 0 else DARK)


# ---------------------------------------------------------------- PNG
def _f(size, bold=False, italic=False):
    return ImageFont.truetype(F_BOLD if bold else (F_IT if italic else F_REG), int(round(size / 72 * PX)))


def _wrap_runs(d, runs, size, width_px):
    """Découpe des runs en lignes ; renvoie [[(texte, police, couleur), ...], ...]."""
    lines, cur, cur_w = [], [], 0
    for r in runs:
        f = _f(size, r.bold, r.italic)
        for k, word in enumerate(r.text.split(" ")):
            token = (word if k == 0 else " " + word)
            if not token:
                continue
            tw = d.textlength(token, font=f)
            if cur_w + tw > width_px and cur:
                lines.append(cur)
                token = token.lstrip()
                tw = d.textlength(token, font=f)
                cur, cur_w = [], 0
            cur.append((token, f, r.color))
            cur_w += tw
    if cur:
        lines.append(cur)
    return lines


def to_png(slide, path, warnings, index):
    img = PImage.new("RGB", (int(W_IN * PX), int(H_IN * PX)), (255, 255, 255))
    d = ImageDraw.Draw(img)
    items = list(slide.items)
    if slide.title:
        items.insert(0, Text(0.5, 0.3, 12.3, 0.7, [Para([Run(slide.title, True)], 26, space_after=0)]))
    for it in items:
        X, Y = int(it.x * PX), int(it.y * PX)
        if isinstance(it, Img):
            im = PImage.open(it.path).convert("RGB").resize((int(it.w * PX), int(it.h * PX)), PImage.LANCZOS)
            img.paste(im, (X, Y))
        elif isinstance(it, Rect):
            d.rectangle([X, Y, X + it.w * PX, Y + it.h * PX], fill=it.fill, outline=it.line, width=2 if it.line else 0)
        elif isinstance(it, Text):
            if it.fill or it.line:
                d.rectangle([X, Y, X + it.w * PX, Y + it.h * PX], fill=it.fill, outline=it.line, width=2)
            yy = Y + 0.05 * PX
            inner = (it.w - 0.16) * PX
            blocks = []
            for p in it.paras:
                runs = ([Run("• ")] + p.runs) if p.bullet else p.runs
                lines = _wrap_runs(d, runs, p.size, inner)
                lh = p.size / 72 * PX * 1.2
                blocks.append((lines, lh, p.space_after / 72 * PX))
            total = sum(len(l) * lh + sa for l, lh, sa in blocks)
            if it.anchor_middle:
                yy = Y + (it.h * PX - total) / 2
            for lines, lh, sa in blocks:
                for line in lines:
                    xx = X + 0.08 * PX
                    for token, f, color in line:
                        d.text((xx, yy), token, font=f, fill=color)
                        xx += d.textlength(token, font=f)
                    yy += lh
                yy += sa
            if yy - (Y + it.h * PX) > 4:
                warnings.append(f"page {index} : texte qui déborde de {((yy - Y) / PX - it.h):.2f} in ({it.paras[0].runs[0].text[:40]!r})")
        elif isinstance(it, Table):
            yy = Y
            for i, row in enumerate(it.rows):
                xx = X
                f = _f(it.size, bold=i == 0)
                rh = it.row_h * PX
                wrapped = []
                for j, val in enumerate(row):
                    lines = _wrap_runs(d, [Run(val, i == 0)], it.size, it.widths[j] * PX - 0.12 * PX)
                    wrapped.append(lines)
                need = max(len(l) for l in wrapped) * it.size / 72 * PX * 1.2 + 0.04 * PX
                rh = max(rh, need)
                for j, lines in enumerate(wrapped):
                    w = it.widths[j] * PX
                    fill = BLUE if i == 0 else ((255, 255, 255) if i % 2 else LIGHT)
                    d.rectangle([xx, yy, xx + w, yy + rh], fill=fill, outline=(200, 200, 200))
                    ty = yy + 0.02 * PX
                    for line in lines:
                        tx = xx + 0.06 * PX
                        for token, ff, _c in line:
                            d.text((tx, ty), token, font=ff, fill=(255, 255, 255) if i == 0 else DARK)
                            tx += d.textlength(token, font=ff)
                        ty += it.size / 72 * PX * 1.2
                    xx += w
                yy += rh
            if yy > H_IN * PX - 0.1 * PX:
                warnings.append(f"page {index} : tableau qui sort de la page")
    img.save(path)
    return img


def to_pdf(slides, pdf_path, png_dir):
    warnings, pages = [], []
    for i, s in enumerate(slides, start=1):
        pages.append(to_png(s, png_dir / f"page-{i:02d}.png", warnings, i))
    pages[0].save(pdf_path, save_all=True, append_images=pages[1:], resolution=PX)
    return warnings
